"""Membuat dokumen uji dengan tata letak beragam (dijalankan sekali; hasilnya disimpan di tests/fixtures).

    python tests/make_fixtures.py

Butuh: reportlab, python-docx, openpyxl, Pillow (hanya untuk membuat ulang berkas uji).
"""
from __future__ import annotations

import random
from datetime import datetime
from pathlib import Path

OUT = Path(__file__).parent / "fixtures"
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
MONTHS = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]


def pdf_invoice():
    """Faktur BBM: tabel item dengan kolom uang; volume di kolom Qty + kolom Satuan terpisah."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(OUT / "faktur_bbm_tabel.pdf"), pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 14); c.drawString(40, h - 50, "PT Energi Nusantara Mandiri")
    c.setFont("Helvetica", 9); c.drawString(40, h - 64, "Jl. Raya Industri No. 45, Dumai · Telp 0765-123456")
    c.setFont("Helvetica-Bold", 16); c.drawRightString(w - 40, h - 50, "FAKTUR")
    c.setFont("Helvetica", 9)
    c.drawRightString(w - 40, h - 66, "No. Faktur: INV/BBM/2025/10/0921")
    c.drawRightString(w - 40, h - 78, "Tanggal: 14-10-2025")
    c.drawString(40, h - 110, "Kepada: PKS Sei Lestari, Kab. Siak")
    y = h - 150
    cols = [40, 70, 250, 320, 380, 470]
    heads = ["No", "Uraian", "Qty", "Satuan", "Harga Satuan (Rp)", "Jumlah (Rp)"]
    c.setFont("Helvetica-Bold", 9)
    for x, t in zip(cols, heads):
        c.drawString(x, y, t)
    c.line(40, y - 4, w - 40, y - 4)
    rows = [("1", "Biosolar B40 industri", "12.000", "Liter", "12.500", "150.000.000"),
            ("2", "Biosolar B40 industri", "4.000", "Liter", "12.500", "50.000.000"),
            ("3", "Ongkos angkut", "1", "Rit", "2.500.000", "2.500.000")]
    c.setFont("Helvetica", 9)
    for r in rows:
        y -= 16
        for i, (x, t) in enumerate(zip(cols, r)):
            if i in (2, 4, 5):
                c.drawRightString(x + 50 if i == 2 else x + 75, y, t)
            else:
                c.drawString(x, y, t)
    y -= 30
    c.drawString(380, y, "Subtotal"); c.drawRightString(w - 40, y, "202.500.000")
    c.drawString(380, y - 14, "PPN 11%"); c.drawRightString(w - 40, y - 14, "22.275.000")
    c.setFont("Helvetica-Bold", 9)
    c.drawString(380, y - 28, "Total"); c.drawRightString(w - 40, y - 28, "Rp 224.775.000")
    c.setFont("Helvetica", 8)
    c.drawString(40, 60, "Keterangan: untuk genset pabrik. Pembayaran ke rekening 123-456-7890.")
    c.save()


def pdf_pln_form():
    """Tagihan listrik bergaya formulir: label di atas nilai, LWBP/WBP terpisah, total di kanan bawah."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(OUT / "tagihan_pln_formulir.pdf"), pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 13); c.drawString(40, h - 50, "TAGIHAN LISTRIK PELANGGAN TEGANGAN MENENGAH")
    c.setFont("Helvetica", 8)
    boxes = [(40, "ID PELANGGAN", "512345678901"), (210, "PERIODE", "AGUSTUS 2025"), (380, "TARIF / DAYA", "I-3 / 1.385.000 VA")]
    for x, lab, val in boxes:
        c.rect(x, h - 120, 160, 46)
        c.setFont("Helvetica", 8); c.drawString(x + 8, h - 88, lab)
        c.setFont("Helvetica-Bold", 11); c.drawString(x + 8, h - 108, val)
    c.setFont("Helvetica", 9)
    c.drawString(40, h - 160, "Stand Meter Awal      120.000")
    c.drawString(240, h - 160, "Stand Meter Akhir      318.400")
    c.drawString(40, h - 180, "Pemakaian LWBP")
    c.drawRightString(230, h - 180, "150.000 kWh")
    c.drawString(40, h - 196, "Pemakaian WBP")
    c.drawRightString(230, h - 196, "48.400 kWh")
    c.drawString(40, h - 212, "Pemakaian kVArh")
    c.drawRightString(230, h - 212, "2.150 kVArh")
    c.rect(360, h - 260, 195, 60)
    c.setFont("Helvetica", 8); c.drawString(370, h - 215, "TOTAL PEMAKAIAN")
    c.setFont("Helvetica-Bold", 16); c.drawString(370, h - 245, "198.400 kWh")
    c.setFont("Helvetica", 9)
    c.drawString(40, h - 300, "Tagihan bulan ini      Rp 287.680.000")
    c.drawString(40, h - 316, "Tgl jatuh tempo: 20-09-2025")
    c.save()


def pdf_production_two_pages():
    """Rekap produksi bulanan; tabel terpotong ke halaman 2, judul diulang, total di akhir."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(OUT / "rekap_produksi_2_halaman.pdf"), pagesize=A4)
    w, h = A4
    tbs = [9000, 9500, 10000, 10200, 10400, 10600, 10800, 10500, 10000, 9800, 9600, 9600]
    heads = ["Bulan", "TBS Diolah (ton)", "CPO (ton)", "PKO (ton)", "Rendemen CPO (%)"]
    xs = [40, 170, 290, 390, 480]

    def header(y):
        c.setFont("Helvetica-Bold", 9)
        for x, t in zip(xs, heads):
            c.drawString(x, y, t)
        c.setFont("Helvetica", 9)

    c.setFont("Helvetica-Bold", 13); c.drawString(40, h - 50, "LAPORAN PRODUKSI PABRIK KELAPA SAWIT 2025")
    c.setFont("Helvetica", 9); c.drawString(40, h - 66, "PKS Sei Lestari — kapasitas 45 ton TBS/jam")
    y = h - 110
    header(y)
    for i in range(12):
        if i == 7:
            c.drawString(40, 40, "Bersambung ke halaman 2")
            c.showPage()
            y = h - 60
            header(y)
        y -= 18
        cpo, pko = tbs[i] * 0.22, tbs[i] * 0.022
        c.drawString(40, y, MONTHS[i])
        c.drawRightString(xs[1] + 60, y, f"{tbs[i]:,}".replace(",", "."))
        c.drawRightString(xs[2] + 50, y, f"{cpo:,.0f}".replace(",", "."))
        c.drawRightString(xs[3] + 40, y, f"{pko:,.1f}".replace(",", "X").replace(".", ",").replace("X", "."))
        c.drawRightString(xs[4] + 50, y, "22,0")
    y -= 22
    c.setFont("Helvetica-Bold", 9)
    c.drawString(40, y, "Total")
    c.drawRightString(xs[1] + 60, y, "120.000")
    c.drawRightString(xs[2] + 50, y, "26.400")
    c.drawRightString(xs[3] + 40, y, "2.640,0")
    c.save()


def pdf_big_number_caption():
    """Angka besar di atas, keterangan di bawahnya; tahun & tanggal tidak boleh terbaca sebagai volume."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(OUT / "ringkasan_solar_angka_atas.pdf"), pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 12); c.drawString(40, h - 50, "RINGKASAN PEMAKAIAN BAHAN BAKAR TAHUN 2025")
    c.setFont("Helvetica", 9); c.drawString(40, h - 66, "Dicetak 05-01-2026 · Divisi Teknik")
    c.setFont("Helvetica-Bold", 28); c.drawString(40, h - 140, "1.533.400 L")
    c.setFont("Helvetica", 10); c.drawString(40, h - 160, "Total pembelian biosolar B40 selama 2025")
    c.setFont("Helvetica", 9); c.drawString(40, h - 200, "Harga rata-rata Rp 12.500 per liter. Stok akhir tahun tercatat di gudang.")
    c.save()


def xlsx_monthly_electricity():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Listrik 2025"
    ws["A1"] = "REKAP PEMAKAIAN LISTRIK PLN"
    ws["A2"] = "PT Sawit Lestari Nusantara · ID Pelanggan 512345678901"
    ws.append([])
    ws.append([])
    ws.append(["No", "Bulan", "Pemakaian (kWh)", "Biaya (Rp)", "Keterangan"])
    for m in range(1, 13):
        ws.append([m, datetime(2025, m, 1), 200000, 290000000, "lunas"])
    ws.append(["", "Total", "=SUM(C6:C17)", "=SUM(D6:D17)"])
    ws["C18"] = 2400000  # nilai tersimpan (openpyxl tidak menghitung rumus)
    ws["D18"] = 3480000000
    for row in ws.iter_rows(min_row=6, max_row=17, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "mmm-yy"
    wb.save(OUT / "listrik_bulanan.xlsx")


def xlsx_keyvalue_pome():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Lab POME"
    ws["B2"] = "LAPORAN HASIL UJI LIMBAH CAIR (POME)"
    ws["B4"] = "Periode"
    ws["D4"] = "Triwulan III 2025"
    ws["B6"] = "Volume POME"
    ws["D6"] = 18000
    ws["E6"] = "m3"
    ws["B8"] = "Parameter"
    ws["C8"] = "Satuan"
    ws["D8"] = "Inlet"
    ws["E8"] = "Outlet"
    ws["F8"] = "Baku mutu"
    ws.append([])
    ws["B9"] = "pH"; ws["D9"] = 4.2; ws["E9"] = 7.6; ws["F9"] = "6-9"
    ws["B10"] = "COD"; ws["C10"] = "mg/L"; ws["D10"] = 51200; ws["E10"] = 4800; ws["F10"] = 350
    ws["B11"] = "BOD"; ws["C11"] = "mg/L"; ws["D11"] = 25000; ws["E11"] = 900; ws["F11"] = 100
    ws["B13"] = "Biogas tertangkap"
    ws["D13"] = 0
    ws["E13"] = "m3"
    wb.save(OUT / "lab_pome_label_kiri.xlsx")


def csv_fertilizer():
    rows = [
        "REKAP APLIKASI PUPUK SEMESTER II 2025;;;",
        ";;;",
        "Afdeling;Jenis Pupuk;Jumlah (kg);Kadar N (%)",
        "I;NPK 15-15-6-4;42.500;15",
        "II;Urea;18.000;46",
        "III;MOP;10.000;0",
        "IV;Dolomit;25.000;",
        "Total;;95.500;",
    ]
    (OUT / "rekap_pupuk_titik_koma.csv").write_text("\n".join(rows), encoding="utf-8")


def docx_b3():
    import docx

    d = docx.Document()
    d.add_heading("MANIFEST LIMBAH B3", 1)
    d.add_paragraph("Nomor manifest: MNF/B3/2025/044")
    d.add_paragraph("Penghasil: PKS Sei Lestari")
    t = d.add_table(rows=1, cols=4)
    for i, hd in enumerate(["Jenis Limbah", "Kode Limbah", "Jumlah (ton)", "Pengolahan"]):
        t.rows[0].cells[i].text = hd
    for r in [("Oli bekas", "B105d", "12", "Insinerasi pihak ketiga"), ("Aki bekas", "A102d", "6,5", "Pemanfaatan oleh pihak ketiga")]:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    d.add_paragraph("Total limbah dikirim: 18,5 ton")
    d.save(OUT / "manifest_b3.docx")


def _render_scan(lines: list[tuple[int, int, str, int, bool]], size=(1240, 1754), rotate=0.0, noise=True):
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    img = Image.new("L", size, 250)
    dr = ImageDraw.Draw(img)
    for x, y, text, pt, bold in lines:
        dr.text((x, y), text, fill=25, font=ImageFont.truetype(FONT_B if bold else FONT, pt))
    if noise:
        rnd = random.Random(7)
        px = img.load()
        for _ in range(9000):
            px[rnd.randrange(size[0]), rnd.randrange(size[1])] = rnd.choice((180, 210, 90))
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    if rotate:
        img = img.rotate(rotate, expand=True, fillcolor=245, resample=3)
    return img


def scan_pome_png():
    """Foto/scan miring 2,5°: label di kiri atas, nilai di kanan bawah, tabel matriks COD."""
    lines = [
        (90, 90, "LABORATORIUM LINGKUNGAN TIRTA", 34, True),
        (90, 145, "Laporan hasil uji air limbah POME", 24, False),
        (90, 230, "Volume POME bulan ini", 26, False),
        (760, 230, "18.000 m3", 30, True),
        (90, 340, "Parameter", 24, True), (420, 340, "Inlet", 24, True), (640, 340, "Outlet", 24, True),
        (90, 395, "COD", 24, False), (420, 395, "51.200", 24, False), (640, 395, "4.800", 24, False),
        (90, 450, "BOD", 24, False), (420, 450, "25.000", 24, False), (640, 450, "900", 24, False),
        (90, 560, "Satuan mg/L. Dicetak 03-10-2025.", 20, False),
    ]
    _render_scan(lines, rotate=2.5).convert("RGB").save(OUT / "scan_pome_miring.jpg", quality=80)


def scan_fuel_pdf_rotated():
    """PDF berisi gambar (tanpa lapisan teks), diputar 90°: perlu OSD + OCR."""
    lines = [
        (90, 90, "BUKTI PENERIMAAN BBM", 34, True),
        (90, 150, "Depot Pertamina Dumai", 24, False),
        (90, 240, "Jenis BBM", 26, False), (560, 240, "Biosolar B40", 26, True),
        (90, 300, "Volume diterima", 26, False), (560, 300, "8.000 Liter", 26, True),
        (90, 360, "Tujuan", 26, False), (560, 360, "Genset pabrik", 26, False),
        (90, 470, "Petugas gudang: Andi · Tanggal 21-07-2025", 20, False),
    ]
    img = _render_scan(lines, rotate=90).convert("RGB")
    img.save(OUT / "scan_bbm_diputar.pdf", "PDF", resolution=150)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for fn in (pdf_invoice, pdf_pln_form, pdf_production_two_pages, pdf_big_number_caption, xlsx_monthly_electricity,
               xlsx_keyvalue_pome, csv_fertilizer, docx_b3, scan_pome_png, scan_fuel_pdf_rotated):
        fn()
        print("ok", fn.__name__)
