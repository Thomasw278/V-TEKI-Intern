"""Ekspor laporan: XLSX (lengkap dengan jejak data) dan XML gaya komunikasi emisi CBAM."""
from __future__ import annotations

import io
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..models import Calculation, Company, Document, Product, ReportingPeriod

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="13432C")


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    ws.freeze_panes = "A2"
    return ws


def _source_label(source: str | None) -> str:
    source = source or "default"
    if source.startswith("profile:"):
        return "Profil kebutuhan"
    if source.startswith("exception:"):
        return "Permintaan pengecualian disetujui"
    return {"default": "Default katalog", "manual": "Diatur superadmin"}.get(source, source)


def build_xlsx(db: Session, period: ReportingPeriod, calc: Calculation) -> bytes:
    company = db.get(Company, period.company_id)
    products = {p.code: p for p in db.scalars(select(Product).where(Product.company_id == company.id))}
    res = calc.result
    excluded = res.get("excluded", [])
    wb = Workbook()
    status = "TERKUNCI" if period.status == "locked" else "DRAF — bukan untuk pelaporan resmi"
    _sheet(wb, "Ringkasan", ["Item", "Nilai"], [
        ["Perusahaan", company.name], ["Alamat", ", ".join(x for x in [company.address, company.city, company.province] if x)],
        ["Periode", period.year], ["Status", status], ["Versi kalkulasi", calc.version],
        ["Set GWP", calc.gwp_set], ["Metode alokasi", "Ekonomi" if calc.allocation_method == "economic" else "Massa"],
        ["Emisi langsung (t CO2e)", float(calc.direct_t)], ["Emisi tidak langsung (t CO2e)", float(calc.indirect_t)],
        ["Emisi hulu (t CO2e)", float(calc.upstream_t)], ["Total (t CO2e)", float(calc.total_t)],
        ["CO2 biogenik dilaporkan terpisah (t)", res.get("biogenic_co2_t", 0)],
        ["Kategori memakai nilai bawaan", ", ".join(res.get("defaults_used", {}).keys()) or "-"],
        ["Kategori dikecualikan (tidak berlaku)", ", ".join(x["name"] for x in excluded) or "-"],
    ], first=True)
    _sheet(wb, "SEE per produk", ["Kode", "Produk", "Kode CN", "Tonase (t)", "Share", "SEE langsung", "SEE tidak langsung", "SEE total (t CO2e/t)"], [
        [a["code"], a["name"], (products.get(a["code"]).cn_code if products.get(a["code"]) else ""),
         a["tonnes"], a["share"], round(a["see_direct"], 6), round(a["see_indirect"], 6), round(a["see_total"], 6)]
        for a in res.get("allocations", [])
    ])
    _sheet(wb, "Rincian komponen", ["Komponen", "Kelompok", "Scope", "Data aktivitas", "Faktor", "Rumus", "Perhitungan", "t CO2e", "Sumber data"], [
        [ln["name"], ln["group"], ln["scope"], ln["activity"], ln["factor"], ln["formula"], ln["calc"], round(ln["value_t"], 3),
         "Nilai bawaan" if ln["source"] == "default" else "Dokumen"]
        for ln in res.get("lines", [])
    ])
    # "Tidak berlaku" ≠ "data tidak tersedia": kategori ini dikeluarkan dari perhitungan, bukan diisi 0
    _sheet(wb, "Kategori dikecualikan", ["Kode", "Kategori", "Alasan", "Ditetapkan oleh", "Waktu", "Asal penetapan"], [
        [x["code"], x["name"], x.get("reason") or "-", x.get("set_by") or "Default katalog",
         (x.get("set_at") or "")[:16].replace("T", " "), _source_label(x.get("source"))]
        for x in excluded
    ] or [["-", "Tidak ada kategori yang dikecualikan", "", "", "", ""]])
    _sheet(wb, "Faktor emisi", ["Kode", "Nilai", "Satuan", "Status", "Versi", "Sumber"], [
        [k, v["value"], v["unit"], v["status"], v["version"], v["source"]] for k, v in (calc.factors_snapshot or {}).items()
    ])
    docs = db.scalars(select(Document).where(Document.period_id == period.id).options(selectinload(Document.fields))
                      .order_by(Document.type_code, Document.id)).all()
    rows = []
    for d in docs:
        for f in d.fields:
            rows.append([d.id, d.original_name, d.type_code or "-", d.status, f.label,
                         f.value_text, f.unit or "", float(f.confidence or 0), "ya" if f.confirmed else "tidak", f.page or "", (f.snippet or "")[:200]])
    _sheet(wb, "Data aktivitas", ["ID dok", "Berkas", "Kategori", "Status", "Field", "Nilai", "Satuan", "Keyakinan", "Dikonfirmasi", "Hal.", "Kutipan sumber"], rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_xml(db: Session, period: ReportingPeriod, calc: Calculation) -> bytes:
    """Struktur sederhana mengikuti isi template komunikasi emisi CBAM.

    Catatan: ini BUKAN skema XSD resmi registri CBAM; sesuaikan dengan
    format yang diminta importir/registri sebelum dikirim.
    """
    company = db.get(Company, period.company_id)
    products = {p.code: p for p in db.scalars(select(Product).where(Product.company_id == company.id))}
    root = ET.Element("EmissionsCommunication", {"version": "1.0", "status": period.status})
    inst = ET.SubElement(root, "Installation")
    for tag, val in [("OperatorName", company.name), ("Address", company.address or ""), ("City", company.city or ""),
                     ("Region", company.province or ""), ("Country", "ID"), ("PostalCode", company.postal_code or "")]:
        ET.SubElement(inst, tag).text = val
    rp = ET.SubElement(root, "ReportingPeriod")
    ET.SubElement(rp, "Start").text = f"{period.year}-01-01"
    ET.SubElement(rp, "End").text = f"{period.year}-12-31"
    meth = ET.SubElement(root, "Methodology")
    ET.SubElement(meth, "GWPSet").text = calc.gwp_set
    ET.SubElement(meth, "AllocationMethod").text = calc.allocation_method
    ET.SubElement(meth, "CalculationVersion").text = str(calc.version)
    goods = ET.SubElement(root, "Goods")
    for a in calc.result.get("allocations", []):
        g = ET.SubElement(goods, "Good")
        p = products.get(a["code"])
        ET.SubElement(g, "ProductCode").text = a["code"]
        ET.SubElement(g, "Name").text = a["name"]
        ET.SubElement(g, "CNCode").text = (p.cn_code if p else "") or ""
        ET.SubElement(g, "ProductionTonnes").text = f"{a['tonnes']:.3f}"
        ET.SubElement(g, "SpecificEmbeddedEmissionsDirect", unit="tCO2e/t").text = f"{a['see_direct']:.6f}"
        ET.SubElement(g, "SpecificEmbeddedEmissionsIndirect", unit="tCO2e/t").text = f"{a['see_indirect']:.6f}"
        ET.SubElement(g, "SpecificEmbeddedEmissionsTotal", unit="tCO2e/t").text = f"{a['see_total']:.6f}"
    tot = ET.SubElement(root, "InstallationTotals", unit="tCO2e")
    ET.SubElement(tot, "Direct").text = f"{float(calc.direct_t):.3f}"
    ET.SubElement(tot, "Indirect").text = f"{float(calc.indirect_t):.3f}"
    ET.SubElement(tot, "Upstream").text = f"{float(calc.upstream_t):.3f}"
    excluded = calc.result.get("excluded", [])
    if excluded:
        ex = ET.SubElement(root, "ExcludedCategories")
        for x in excluded:
            attrs = {"code": x["code"], "name": x["name"]}
            if x.get("set_by"):
                attrs["setBy"] = x["set_by"]
            if x.get("set_at"):
                attrs["setAt"] = x["set_at"]
            ET.SubElement(ex, "Category", attrs).text = x.get("reason") or ""
    defaults = calc.result.get("defaults_used", {})
    if defaults:
        dv = ET.SubElement(root, "DefaultValuesUsed")
        for k, why in defaults.items():
            ET.SubElement(dv, "Category", code=k).text = why
    ET.indent(root)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)
