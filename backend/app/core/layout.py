"""Model tata letak dokumen: kata + koordinat, baris, segmen (kolom), halaman.

Semua format (PDF teks, PDF hasil scan, foto/scan, XLSX, CSV, DOCX, TXT) diubah ke
struktur yang sama sehingga mesin ekstraksi bisa mencari informasi di posisi mana
pun: di kanan label, di bawah label, di tabel, di atas, atau di bagian bawah halaman.

Sistem koordinat per jenis halaman:
- "pdf" / "image": titik atau piksel, origin kiri-atas (x ke kanan, y ke bawah)
- "text": x = posisi karakter, y = nomor baris
- "sheet": x = indeks kolom, y = indeks baris (XLSX, CSV, DOCX)
"""
from __future__ import annotations

import csv
import io
import re
import statistics
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET


@dataclass
class Word:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    conf: float = 1.0

    @property
    def cx(self) -> float:
        return (self.x0 + self.x1) / 2

    @property
    def cy(self) -> float:
        return (self.y0 + self.y1) / 2

    @property
    def h(self) -> float:
        return max(self.y1 - self.y0, 1e-6)


@dataclass
class Segment:
    """Sekelompok kata yang berdekatan di satu baris (≈ satu sel/kolom)."""
    words: list[Word]

    @property
    def text(self) -> str:
        return " ".join(w.text for w in self.words)

    @property
    def x0(self) -> float:
        return min(w.x0 for w in self.words)

    @property
    def x1(self) -> float:
        return max(w.x1 for w in self.words)

    @property
    def cx(self) -> float:
        return (self.x0 + self.x1) / 2


@dataclass
class Line:
    page: int
    index: int
    segments: list[Segment]

    @property
    def words(self) -> list[Word]:
        return [w for s in self.segments for w in s.words]

    @property
    def text(self) -> str:
        return "  ".join(s.text for s in self.segments)

    @property
    def x0(self) -> float:
        return min(w.x0 for w in self.words)

    @property
    def x1(self) -> float:
        return max(w.x1 for w in self.words)

    @property
    def y0(self) -> float:
        return min(w.y0 for w in self.words)

    @property
    def y1(self) -> float:
        return max(w.y1 for w in self.words)


@dataclass
class Page:
    number: int
    kind: str  # pdf | image | text | sheet
    width: float
    height: float
    lines: list[Line] = field(default_factory=list)
    name: str = ""
    ocr: bool = False
    preview_lines: list[str] | None = None          # untuk "text"
    preview_rows: list[list[str]] | None = None     # untuk "sheet"


@dataclass
class DocLayout:
    pages: list[Page]
    method: str
    warnings: list[str] = field(default_factory=list)

    @property
    def lines(self) -> list[Line]:
        return [ln for p in self.pages for ln in p.lines]

    @property
    def text(self) -> str:
        return "\n\f\n".join("\n".join(ln.text for ln in p.lines) for p in self.pages)

    def page(self, n: int) -> Page:
        return self.pages[n]

    def bbox(self, page_no: int, words: list[Word]) -> dict | None:
        """Kotak sumber yang dinormalisasi untuk ditampilkan di pratinjau."""
        if not words:
            return None
        p = self.pages[page_no]
        x0, y0 = min(w.x0 for w in words), min(w.y0 for w in words)
        x1, y1 = max(w.x1 for w in words), max(w.y1 for w in words)
        if p.kind in ("text", "sheet"):
            return {"page": page_no, "r0": int(y0), "r1": int(y1) - 1 if y1 > y0 + 0.5 else int(y0),
                    "c0": int(x0), "c1": max(int(x1) - 1, int(x0))}
        pad = 2.0
        return {
            "page": page_no,
            "x0": round(max(0.0, (x0 - pad) / p.width), 4),
            "y0": round(max(0.0, (y0 - pad) / p.height), 4),
            "x1": round(min(1.0, (x1 + pad) / p.width), 4),
            "y1": round(min(1.0, (y1 + pad) / p.height), 4),
        }

    def preview(self) -> dict:
        """Struktur ringan untuk pratinjau di frontend (disimpan di database)."""
        out = []
        for p in self.pages:
            d: dict = {"n": p.number, "kind": p.kind, "w": round(p.width, 1), "h": round(p.height, 1), "ocr": p.ocr}
            if p.name:
                d["name"] = p.name
            if p.kind == "text":
                d["lines"] = (p.preview_lines or [])[:600]
            elif p.kind == "sheet":
                d["rows"] = [r[:40] for r in (p.preview_rows or [])[:400]]
            out.append(d)
        return {"method": self.method, "pages": out}


# ---------------------------------------------------------------- utilitas baris
def group_lines(words: list[Word], page_no: int, gap_factor: float = 1.6) -> list[Line]:
    """Kelompokkan kata menjadi baris visual, lalu pecah baris menjadi segmen (kolom)."""
    words = [w for w in words if w.text.strip()]
    if not words:
        return []
    words.sort(key=lambda w: (w.cy, w.x0))
    med_h = statistics.median(w.h for w in words)
    rows: list[list[Word]] = []
    centers: list[float] = []
    for w in words:
        placed = False
        # bandingkan dengan beberapa baris terakhir (dokumen miring sedikit tetap tertangkap)
        for i in range(len(rows) - 1, max(-1, len(rows) - 4), -1):
            if abs(centers[i] - w.cy) <= med_h * 0.55:
                rows[i].append(w)
                centers[i] = statistics.fmean(x.cy for x in rows[i])
                placed = True
                break
        if not placed:
            rows.append([w])
            centers.append(w.cy)
    order = sorted(range(len(rows)), key=lambda i: centers[i])
    lines = []
    for idx, i in enumerate(order):
        ws = sorted(rows[i], key=lambda w: w.x0)
        lines.append(Line(page_no, idx, split_segments(ws, gap_factor)))
    return lines


def split_segments(ws: list[Word], gap_factor: float = 1.6) -> list[Segment]:
    if not ws:
        return []
    char_w = statistics.median(max((w.x1 - w.x0) / max(len(w.text), 1), 0.5) for w in ws)
    space = char_w * 1.0
    segs: list[list[Word]] = [[ws[0]]]
    for prev, w in zip(ws, ws[1:]):
        gap = w.x0 - prev.x1
        if gap > max(space * gap_factor * 2.2, char_w * 3.2):
            segs.append([w])
        else:
            segs[-1].append(w)
    return [Segment(s) for s in segs]


# ---------------------------------------------------------------- TXT / MD
_SEG_SPLIT = re.compile(r"\t+|\s{2,}|\s*\|\s*")


def text_page(text: str, page_no: int = 0) -> Page:
    raw_lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    lines: list[Line] = []
    for y, raw in enumerate(raw_lines):
        raw = raw.expandtabs(4)
        if not raw.strip():
            continue
        segs = []
        pos = 0
        for part in _SEG_SPLIT.split(raw):
            if not part:
                continue
            start = raw.find(part, pos)
            pos = start + len(part)
            words = []
            for m in re.finditer(r"\S+", part):
                words.append(Word(m.group(0), start + m.start(), y, start + m.end(), y + 1))
            if words:
                segs.append(Segment(words))
        if segs:
            lines.append(Line(page_no, len(lines), segs))
    width = max((len(r) for r in raw_lines), default=1)
    return Page(page_no, "text", float(width), float(len(raw_lines)), lines, preview_lines=raw_lines)


# ---------------------------------------------------------------- grid (XLSX, CSV, DOCX)
def _cell_text(v, number_format: str = "") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "ya" if v else "tidak"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        if "%" in (number_format or ""):
            return f"{v * 100:g}%"
        if v.is_integer():
            return str(int(v))
        return f"{v:.6f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        if "%" in (number_format or ""):
            return f"{v * 100}%"
        return str(v)
    return str(v).strip()


def grid_page(rows: list[list[str]], page_no: int, name: str = "") -> Page:
    lines: list[Line] = []
    for y, row in enumerate(rows):
        segs = []
        for x, val in enumerate(row):
            val = (val or "").strip()
            if not val:
                continue
            words = [Word(t, x + i / max(len(val.split()), 1) * 0.9, y, x + (i + 1) / max(len(val.split()), 1) * 0.9, y + 1)
                     for i, t in enumerate(val.split())]
            # seluruh sel = satu segmen; posisi kolom = indeks sel
            for w in words:
                w.x0, w.x1 = float(x), float(x + 1)
            segs.append(Segment(words))
        if segs:
            lines.append(Line(page_no, len(lines), segs))
    width = max((len(r) for r in rows), default=1)
    return Page(page_no, "sheet", float(width), float(len(rows)), lines, name=name, preview_rows=rows)


def xlsx_layout(path: Path) -> DocLayout:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets):
        rows: list[list[str]] = []
        for r_idx, row in enumerate(ws.iter_rows()):
            if r_idx >= 3000:
                break
            rows.append([_cell_text(c.value, getattr(c, "number_format", "") or "") for c in row[:60]])
        while rows and not any(rows[-1]):
            rows.pop()
        if any(any(c for c in r) for r in rows):
            pages.append(grid_page(rows, len(pages), ws.title))
    return DocLayout(pages, "xlsx")


def csv_layout(path: Path) -> DocLayout:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    text = text or ""
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if text.count(";") > text.count(",") else ","
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)][:3000]
    return DocLayout([grid_page(rows, 0, "CSV")], "csv")


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_layout(path: Path) -> DocLayout:
    """DOCX: paragraf → baris 1 sel, tabel → baris multi-sel (urutan dokumen dipertahankan)."""
    with zipfile.ZipFile(str(path)) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(f"{W_NS}body")
    rows: list[list[str]] = []

    def para_text(p) -> str:
        return "".join(t.text or "" for t in p.iter(f"{W_NS}t")).strip()

    for el in list(body) if body is not None else []:
        if el.tag == f"{W_NS}p":
            t = para_text(el)
            if t:
                rows.append([t])
        elif el.tag == f"{W_NS}tbl":
            for tr in el.iter(f"{W_NS}tr"):
                cells = [" ".join(para_text(p) for p in tc.iter(f"{W_NS}p")).strip() for tc in tr.findall(f"{W_NS}tc")]
                if any(cells):
                    rows.append(cells)
            rows.append([])
    return DocLayout([grid_page(rows, 0, "Dokumen")], "docx")


# ---------------------------------------------------------------- PDF
def pdf_layout(path: Path, ocr_cache_dir: Path | None = None) -> DocLayout:
    """PDF teks dibaca langsung; halaman tanpa teks (hasil scan) di-OCR per halaman."""
    import pdfplumber

    from . import ocr

    pages: list[Page] = []
    warnings: list[str] = []
    used_ocr = False
    with pdfplumber.open(str(path)) as pdf:
        for i, pg in enumerate(pdf.pages[:200]):
            words = pg.extract_words(x_tolerance=1.5, y_tolerance=2.5, keep_blank_chars=False, use_text_flow=False)
            text_chars = sum(len(w["text"]) for w in words)
            if text_chars >= 25:
                ws = [Word(w["text"], float(w["x0"]), float(w["top"]), float(w["x1"]), float(w["bottom"])) for w in words]
                pages.append(Page(i, "pdf", float(pg.width), float(pg.height), group_lines(ws, i)))
                continue
            # halaman hasil scan → OCR
            img = pg.to_image(resolution=300).original
            res = ocr.ocr_image(img)
            used_ocr = True
            if res is None:
                warnings.append(f"Halaman {i + 1} berupa gambar tetapi OCR (Tesseract) belum terpasang.")
                pages.append(Page(i, "pdf", float(pg.width), float(pg.height), []))
                continue
            if ocr_cache_dir:
                res.image.save(ocr_cache_dir / f"p{i}.png", optimize=True)
            pages.append(Page(i, "image", float(res.width), float(res.height), group_lines(res.words, i), ocr=True))
            warnings.extend(f"Halaman {i + 1}: {w}" for w in res.notes)
    return DocLayout(pages, "pdf+ocr" if used_ocr else "pdf", warnings)


def image_layout(path: Path, ocr_cache_dir: Path | None = None) -> DocLayout:
    from PIL import Image

    from . import ocr

    pages: list[Page] = []
    warnings: list[str] = []
    with Image.open(str(path)) as im:
        n = getattr(im, "n_frames", 1)
        for i in range(min(n, 50)):
            im.seek(i)
            frame = im.copy()
            res = ocr.ocr_image(frame)
            if res is None:
                return DocLayout([Page(0, "image", float(frame.width), float(frame.height), [])], "none",
                                 ["OCR (Tesseract) belum terpasang — isi field secara manual."])
            if ocr_cache_dir:
                res.image.save(ocr_cache_dir / f"p{i}.png", optimize=True)
            pages.append(Page(i, "image", float(res.width), float(res.height), group_lines(res.words, i), ocr=True))
            warnings.extend(res.notes)
    return DocLayout(pages, "ocr", warnings)


IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}


def build_layout(path: str | Path, ocr_cache_dir: str | Path | None = None) -> DocLayout:
    p = Path(path)
    ext = p.suffix.lower()
    cache = Path(ocr_cache_dir) if ocr_cache_dir else None
    if cache:
        cache.mkdir(parents=True, exist_ok=True)
    try:
        if ext == ".pdf":
            return pdf_layout(p, cache)
        if ext in (".xlsx", ".xlsm"):
            return xlsx_layout(p)
        if ext == ".csv":
            return csv_layout(p)
        if ext == ".docx":
            return docx_layout(p)
        if ext in IMAGE_EXT:
            return image_layout(p, cache)
        if ext in (".txt", ".md", ".text"):
            raw = p.read_bytes()
            for enc in ("utf-8-sig", "cp1252", "latin-1"):
                try:
                    return DocLayout([text_page(raw.decode(enc))], "txt")
                except UnicodeDecodeError:
                    continue
    except Exception as exc:  # berkas rusak / terenkripsi
        return DocLayout([], "error", [f"Berkas tidak bisa dibaca: {exc}"])
    return DocLayout([], "none", [f"Format {ext} belum didukung"])


def layout_from_text(text: str) -> DocLayout:
    return DocLayout([text_page(text)], "txt")
